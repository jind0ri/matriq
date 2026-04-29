from openpyxl import load_workbook

from app.routers.sample import build_official_sample_report_excel


def test_official_sample_report_excel_contains_core_fields():
    buffer = build_official_sample_report_excel(
        {
            "sample_id": "MAR-2026-001",
            "client_name": "Acme Builders",
            "project_reference": "PRJ-42",
            "material_type": "Concrete",
            "ai_predicted_label": "Concrete",
            "ai_confidence_score": 0.91,
            "decision": "Auto-Accepted",
            "model_version": "v20260416130223",
            "branch_id": 1,
            "current_state": "Released",
            "device_metadata": {
                "payment": {
                    "payment_status": "Fully Paid",
                    "financially_cleared_for_release": True,
                },
                "qa": {
                    "release_reviewed_by_name": "QA Engineer",
                    "release_reviewed_at": "2026-04-29T10:00:00",
                },
                "test_data": {
                    "test_type": "concrete_compression",
                    "system_result": "PASS",
                    "final_result": "PASS",
                    "values": {
                        "test_name": "Concrete Compression Test",
                        "standard": "ASTM C39",
                        "compressive_strength_mpa": 28.4,
                    },
                },
            },
        }
    )

    workbook = load_workbook(buffer)
    sheet = workbook["Official Report"]
    values = [
        cell
        for row in sheet.iter_rows(values_only=True)
        for cell in row
        if cell is not None
    ]

    assert "OFFICIAL LABORATORY TEST REPORT" in values
    assert "MAR-2026-001" in values
    assert "Acme Builders" in values
    assert "91.0%" in values
    assert "Concrete Compression Test" in values
    assert "Employee Signature" in values
